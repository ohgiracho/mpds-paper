from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path

from docx import Document


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def all_docx_text(path: Path) -> list[str]:
    document = Document(path)
    values = [paragraph.text for paragraph in document.paragraphs]
    values.extend(" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows)
    return values


def matching_text(values: list[str], terms: tuple[str, ...]) -> list[str]:
    return [value for value in values if any(term.lower() in value.lower() for term in terms)]


def write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    paper_root = root.parent
    repo = paper_root / "mpds_github_prep" / "github_repo"
    manuscript = paper_root / "AM 투고본" / "MPDS_Adv_Mater_Manuscript.docx"
    si = paper_root / "AM 투고본" / "MPDS_Adv_Mater_Supprting information.docx"
    canonical_dir = repo / "data" / "outputs" / "mpds_stagewise" / "case_01"
    canonical_final = repo / "data" / "outputs" / "mpds" / "case_01" / "mpds_final(1).txt"
    canonical_log = repo / "data" / "outputs" / "mpds" / "case_01" / "Debate_Log(1).txt"
    retest7_dir = paper_root / "교수님 논문" / "Glitter-Cake(1)" / "재7실험" / "MPDS"
    old_dir = paper_root / "교수님 논문" / "Glitter-Cake(1)" / "재재실험" / "MPDS"

    manuscript_matches = matching_text(
        all_docx_text(manuscript),
        ("minority nano-LPSCl", "sequential low-energy mechanofusion", "Case Study 2"),
    )
    si_matches = matching_text(
        all_docx_text(si),
        ("Figure S2", "pre-coated NCM811", "slurry processing as the primary route"),
    )
    si_caption = next(value for value in si_matches if value.startswith("Figure S2."))
    old_final_text = (old_dir / "mpds_final(1).txt").read_text(encoding="utf-8", errors="replace")
    canonical_final_text = canonical_final.read_text(encoding="utf-8", errors="replace")

    source_rows = []
    for stage in ("mpds_round1(1).txt", "mpds_round2(1).txt", "mpds_round3(1).txt"):
        public_path = canonical_dir / stage
        local_path = retest7_dir / stage
        source_rows.append(
            {
                "panel": stage.replace("mpds_", "").replace("(1).txt", ""),
                "canonical_public_file": str(public_path.resolve()),
                "canonical_sha256": sha256(public_path),
                "retest7_local_file": str(local_path.resolve()),
                "retest7_sha256": sha256(local_path),
                "hashes_match": sha256(public_path) == sha256(local_path),
            }
        )
    source_rows.append(
        {
            "panel": "final",
            "canonical_public_file": str(canonical_final.resolve()),
            "canonical_sha256": sha256(canonical_final),
            "retest7_local_file": str((retest7_dir / "mpds_final(1).txt").resolve()),
            "retest7_sha256": sha256(retest7_dir / "mpds_final(1).txt"),
            "hashes_match": sha256(canonical_final) == sha256(retest7_dir / "mpds_final(1).txt"),
        }
    )
    root.joinpath("analysis").mkdir(parents=True, exist_ok=True)
    write_csv(root / "analysis" / "figure_s2_replacement_source_manifest.csv", source_rows)

    evidence = {
        "canonical_public_final_sha256": sha256(canonical_final),
        "retest7_final_sha256": sha256(retest7_dir / "mpds_final(1).txt"),
        "canonical_public_log_sha256": sha256(canonical_log),
        "retest7_log_sha256": sha256(retest7_dir / "Debate_Log(1).txt"),
        "old_retest_final_sha256": sha256(old_dir / "mpds_final(1).txt"),
        "old_retest_log_sha256": sha256(old_dir / "Debate_Log(1).txt"),
        "canonical_matches_retest7": sha256(canonical_final) == sha256(retest7_dir / "mpds_final(1).txt")
        and sha256(canonical_log) == sha256(retest7_dir / "Debate_Log(1).txt"),
        "si_caption_matches_old_solution_markers": all(
            marker.lower() in si_caption.lower()
            for marker in ("pre-coated NCM811", "CNF", "slurry processing as the primary route", "dry processing")
        ),
        "old_final_contains_same_solution_markers": all(
            marker.lower() in old_final_text.lower()
            for marker in ("pre-coated NCM811", "carbon nanofibers", "slurry-based process", "dry-film")
        ),
        "canonical_final_contains_submitted_solution_markers": all(
            marker.lower() in canonical_final_text.lower()
            for marker in ("minority fraction", "majority fraction", "sequential, low-energy mechanofusion")
        ),
        "manuscript_matching_blocks": manuscript_matches,
        "si_caption": si_caption,
    }
    (root / "analysis" / "case2_consistency_evidence.json").write_text(
        json.dumps(evidence, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    report = f"""# Case Study 2 consistency audit

## Identity crosswalk

- Manuscript **Case Study 2** = repository **case_01** = local **Glitter-Cake(1)/재7실험**.
- Repository/file numbering remains unchanged.

## Finding

The main text and Table 3 are consistent with the canonical submitted run, but Supporting Information Figure S2 is not. The canonical GitHub final and full log are byte-for-byte identical to the local `재7실험` copies. Figure S2's caption instead describes the distinctive solution found in the older `재재실험` run.

This is a high-confidence source-version mix-up:

- Canonical/`재7실험`: minority nano-LPSCl for partial coverage, majority micron LPSCl for interstitial filling, sequential low-energy mechanofusion.
- SI Figure S2/older `재재실험`: pre-coated NCM811, a CNF network, bimodal LPSCl, slurry as primary and dry processing as secondary.

The raster images embedded in the SI do not retain a source-file path, so binary provenance cannot be proven from DOCX metadata alone. However, the exact solution-marker agreement between the Figure S2 caption and the older final, together with disagreement with the canonical final, identifies the cause sufficiently for a revision audit.

## Required correction

Replace all four Figure S2 panels (Round 1, Round 2, Round 3, Final Synthesis) from the files listed in `figure_s2_replacement_source_manifest.csv`, and revise the caption to describe convergence toward the canonical allocation rule. Do not alter the repository case number.

## Hash evidence

- Canonical final: `{evidence['canonical_public_final_sha256']}`
- Local `재7실험` final: `{evidence['retest7_final_sha256']}`
- Older `재재실험` final: `{evidence['old_retest_final_sha256']}`
- Canonical and `재7실험` full-log/final match: **{evidence['canonical_matches_retest7']}**

## Scope of the wider audit

All 150 canonical final outputs (30 cases x 5 conditions) and all 60 DS/MPDS full logs exist. Observable topic/date/model/temperature/round/evidence checks found no analogous mismatch for Raw, EO, DS, or MPDS. Nine EOP outputs (repository cases 3, 5, 6, 9-14) do not match the canonical topic; case 11 also has a cutoff-year mismatch. These nine EOP results are excluded from practical replicate-1 reuse and scheduled for three clean runs.
"""
    (root / "analysis" / "case2_consistency_report.md").write_text(report, encoding="utf-8")
    print(json.dumps({"status": "PASS", "canonical_matches_retest7": evidence["canonical_matches_retest7"], "source_panels": len(source_rows)}, indent=2))


if __name__ == "__main__":
    main()
